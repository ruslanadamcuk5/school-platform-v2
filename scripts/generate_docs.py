# -*- coding: utf-8 -*-
"""Generate DOCX versions of the project Markdown documentation."""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUT = DOCS / "word"

PROJECT = "ЕООП «Школьник»"
VERSION = "0.7.1"
DATE = "30.10.2026"


def set_cell_shading(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def add_toc(document: Document) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "separate")
    fld_char3 = OxmlElement("w:fldChar")
    fld_char3.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char)
    run._r.append(instr)
    run._r.append(fld_char2)
    run._r.append(fld_char3)
    note = document.add_paragraph(
        "Оглавление обновляется в Word: Ссылки → Обновить таблицу."
    )
    note.runs[0].italic = True
    note.runs[0].font.size = Pt(9)


def style_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(1.5)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(4)
    r = normal.element.rPr.rFonts
    r.set(qn("w:eastAsia"), "Times New Roman")

    for name, size, color in [
        ("Heading 1", 18, RGBColor(0x1F, 0x3A, 0x5F)),
        ("Heading 2", 15, RGBColor(0x1F, 0x3A, 0x5F)),
        ("Heading 3", 13, RGBColor(0x2E, 0x4E, 0x73)),
    ]:
        st = styles[name]
        st.font.name = "Times New Roman"
        st.font.size = Pt(size)
        st.font.color.rgb = color
        st.font.bold = True
        st.paragraph_format.space_before = Pt(10)
        st.paragraph_format.space_after = Pt(5)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run(f"{PROJECT} · {VERSION} · ")
    add_page_number(footer)


def add_title_page(document: Document, title: str, source: Path) -> None:
    for _ in range(3):
        document.add_paragraph()
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(PROJECT)
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(18)

    document.add_paragraph()
    table = document.add_table(rows=5, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    rows = [
        ("Версия", VERSION),
        ("Статус", "Черновик / рабочая версия"),
        ("Дата", DATE),
        ("Исходный файл", str(source.relative_to(ROOT))),
        ("Формат", "DOCX, сгенерирован из Markdown"),
    ]
    for row, (k, v) in zip(table.rows, rows):
        row.cells[0].text = k
        row.cells[1].text = v
        row.cells[0].paragraphs[0].runs[0].bold = True
        set_cell_shading(row.cells[0], "E8EEF5")

    document.add_paragraph()
    note = document.add_paragraph(
        "Документ является проектной версией. Юридически значимые разделы "
        "требуют согласования с заказчиком, юристом и специалистом по ИБ."
    )
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.runs[0].italic = True
    note.runs[0].font.size = Pt(9)

    document.add_page_break()
    add_toc(document)
    document.add_page_break()


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`.+?`|\[.+?\]\(.+?\)|\*[^*]+?\*)")


def add_inline(paragraph, text: str) -> None:
    text = text.replace("✅", "☑").replace("❌", "☐")
    parts = INLINE_RE.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif part.startswith("[") and "](" in part and part.endswith(")"):
            label, url = part[1:-1].split("](", 1)
            run = paragraph.add_run(label)
            run.underline = True
            run.font.color.rgb = RGBColor(0x1F, 0x4F, 0x8A)
            run2 = paragraph.add_run(f" ({url})")
            run2.font.size = Pt(9)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            paragraph.add_run(part)


def add_code_block(document: Document, lines: list[str]) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.5)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F3F3F3")
    ppr.append(shd)


def add_markdown_table(document: Document, lines: list[str]) -> None:
    rows = []
    for line in lines:
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c or "-") for c in cells):
            continue
        rows.append(cells)
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = document.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            add_inline(p, row[j] if j < len(row) else "")
            p.paragraph_format.space_after = Pt(1)
            for run in p.runs:
                run.font.size = Pt(9)
                if i == 0:
                    run.bold = True
            if i == 0:
                set_cell_shading(cell, "E8EEF5")
    document.add_paragraph()


def render_markdown(document: Document, text: str) -> None:
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            fence = stripped[3:].strip()
            code = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            if fence.lower() == "mermaid":
                add_code_block(document, ["[Схема Mermaid — исходник]"] + code)
            else:
                add_code_block(document, code)
            i += 1
            continue

        if stripped.startswith("|") and stripped.endswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_markdown_table(document, table_lines)
            continue

        if stripped == "---":
            p = document.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            pPr = p._p.get_or_add_pPr()
            pbdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), "B7C4D4")
            pbdr.append(bottom)
            pPr.append(pbdr)
            i += 1
            continue

        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text_value = m.group(2).strip()
            if level == 1:
                p = document.add_paragraph(style="Heading 1")
            elif level == 2:
                p = document.add_paragraph(style="Heading 2")
            elif level == 3:
                p = document.add_paragraph(style="Heading 3")
            else:
                p = document.add_paragraph(style="Heading 3")
            add_inline(p, text_value)
            i += 1
            continue

        if stripped.startswith(">"):
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.7)
            p.paragraph_format.right_indent = Cm(0.5)
            add_inline(p, stripped.lstrip("> ").strip())
            for run in p.runs:
                run.italic = True
            i += 1
            continue

        if re.match(r"^[-*]\s+\[[ xX]\]\s*", stripped):
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            checked = "[x]" in stripped.lower()
            text_value = re.sub(r"^[-*]\s+\[[ xX]\]\s*", "", stripped)
            add_inline(p, ("☑ " if checked else "☐ ") + text_value)
            i += 1
            continue

        if re.match(r"^[-*]\s+", stripped):
            p = document.add_paragraph(style="List Bullet")
            add_inline(p, re.sub(r"^[-*]\s+", "", stripped))
            i += 1
            continue

        if re.match(r"^\d+\.\s+", stripped):
            p = document.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\d+\.\s+", "", stripped))
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        p = document.add_paragraph()
        add_inline(p, stripped)
        i += 1


def build_doc(source: Path, output: Path) -> None:
    text = source.read_text(encoding="utf-8")
    title = source.stem
    for line in text.splitlines():
        if line.startswith("# "):
            title = line[2:].strip()
            break
    document = Document()
    style_document(document)
    add_title_page(document, title, source)
    render_markdown(document, text)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    sources = sorted(p for p in DOCS.rglob("*.md") if "word" not in p.parts)
    if not sources:
        print("No Markdown documents found.")
        return 1
    generated = []
    for source in sources:
        rel = source.relative_to(DOCS)
        output = OUT / rel.with_suffix(".docx")
        build_doc(source, output)
        generated.append(output)
        print(f"OK  {rel} -> {output.relative_to(ROOT)}")

    combined = OUT / "ЕООП_Школьник_Полный_пакет_документации_v0.7.1.docx"
    document = Document()
    style_document(document)
    add_title_page(document, "Полный пакет проектной документации", DOCS / "README.md")
    for source in sources:
        if source.name == "README.md" and source.parent == DOCS:
            continue
        text = source.read_text(encoding="utf-8")
        document.add_page_break()
        label = document.add_paragraph()
        label_run = label.add_run(source.relative_to(DOCS).as_posix())
        label_run.italic = True
        label_run.font.size = Pt(9)
        render_markdown(document, text)
    try:
        document.save(combined)
        print(f"OK  combined -> {combined.relative_to(ROOT)}")
    except PermissionError:
        fallback = combined.with_name(f"{combined.stem}_v0.2{combined.suffix}")
        document.save(fallback)
        print(
            "WARN combined file is open/locked; saved -> "
            f"{fallback.relative_to(ROOT)}"
        )
    print(f"Generated {len(generated)} individual DOCX files.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
